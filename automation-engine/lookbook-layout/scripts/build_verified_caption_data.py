#!/usr/bin/env python3
"""Build caption-data.tsv only from a reviewed caption-map.tsv and its Excel source."""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
from pathlib import Path


MAP_FIELDS = [
    "look_id", "excel_sheet", "excel_look_number", "excel_image",
    "left_filename", "right_filename", "evidence_file", "visual_status",
]
CAPTION_FIELDS = ["look_id", "type", "brand", "price", "article"]


def digest(path: Path) -> str:
    value = hashlib.sha256()
    with path.open("rb") as source:
        for block in iter(lambda: source.read(1024 * 1024), b""):
            value.update(block)
    return value.hexdigest()


def clean(value: object) -> str:
    return " ".join(str("" if value is None else value).replace("\xa0", " ").split())


def price(value: object) -> str:
    if isinstance(value, bool) or value in (None, ""):
        return ""
    if isinstance(value, (int, float)) and float(value).is_integer():
        return f"{int(value):,}".replace(",", " ") + " ₽"
    return clean(value)


def read_map(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8-sig") as source:
        reader = csv.DictReader(source, delimiter="\t")
        if reader.fieldnames != MAP_FIELDS:
            raise ValueError("caption-map.tsv has an unexpected column schema.")
        rows = [{key: clean(value) for key, value in row.items()} for row in reader]
    seen: set[tuple[str, str]] = set()
    for index, row in enumerate(rows, start=1):
        expected = f"LOOK_{index:03}"
        if row["look_id"] != expected:
            raise ValueError(f"Caption map row {index + 1}: expected {expected}.")
        if row["visual_status"] != "CONFIRMED":
            raise ValueError(f"{expected}: visual match is not confirmed.")
        key = (row["excel_sheet"], row["excel_look_number"])
        if row["excel_sheet"] not in {"W", "M"} or not row["excel_look_number"].isdigit() or key in seen:
            raise ValueError(f"{expected}: invalid or duplicate Excel mapping.")
        seen.add(key)
    if not rows:
        raise ValueError("Caption map has no rows.")
    return rows


def extract_products(workbook_path: Path, mapping: list[dict[str, str]]) -> list[dict[str, str]]:
    import openpyxl

    workbook = openpyxl.load_workbook(workbook_path, data_only=True, read_only=False)
    grouped: dict[tuple[str, str], list[tuple[str, str, str, str]]] = {}
    try:
        for sheet_name in ("W", "M"):
            if sheet_name not in workbook.sheetnames:
                raise ValueError(f"Workbook is missing worksheet {sheet_name}.")
            current = ""
            for values in workbook[sheet_name].iter_rows(min_row=2, values_only=True):
                number, kind, brand, amount, article = (list(values) + [None] * 6)[1:6]
                number_text = clean(number)
                if number_text.isdigit():
                    current = str(int(number_text))
                fields = (clean(kind), clean(brand), price(amount), clean(article))
                if current and all(fields):
                    grouped.setdefault((sheet_name, current), []).append(fields)
    finally:
        workbook.close()
    captions: list[dict[str, str]] = []
    for row in mapping:
        products = grouped.get((row["excel_sheet"], row["excel_look_number"]), [])
        if not products:
            raise ValueError(f"{row['look_id']}: no complete products in its mapped Excel group.")
        for kind, brand, amount, article in products:
            captions.append({"look_id": row["look_id"], "type": kind, "brand": brand, "price": amount, "article": article})
    return captions


def write_json(path: Path, value: dict[str, object]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(json.dumps(value, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    os.replace(temporary, path)


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate verified lookbook credits from an approved caption map.")
    parser.add_argument("caption_map", type=Path)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("captions", type=Path)
    parser.add_argument("--provenance", type=Path, required=True)
    args = parser.parse_args()
    try:
        mapping = read_map(args.caption_map)
        captions = extract_products(args.workbook, mapping)
    except ValueError as error:
        raise SystemExit(f"BLOCKED: {error}") from error
    args.captions.parent.mkdir(parents=True, exist_ok=True)
    with args.captions.open("w", encoding="utf-8", newline="") as target:
        writer = csv.DictWriter(target, fieldnames=CAPTION_FIELDS, delimiter="\t")
        writer.writeheader()
        writer.writerows(captions)
    write_json(args.provenance, {
        "schema": 1,
        "caption_map_sha256": digest(args.caption_map),
        "caption_workbook_sha256": digest(args.workbook),
        "caption_data_sha256": digest(args.captions),
        "look_count": len(mapping),
        "caption_rows": len(captions),
    })
    print(f"Wrote {len(captions)} verified product rows for {len(mapping)} visually confirmed looks.")


if __name__ == "__main__":
    main()
