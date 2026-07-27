#!/usr/bin/env python3
"""Retired unsafe entry point.

This former tool used Excel references embedded in the layout registry, which
allowed a guessed ordering to become live credits.  It is intentionally blocked;
use build_verified_caption_data.py with a reviewed caption-map.tsv instead.
"""
from __future__ import annotations

import csv
import sys
from pathlib import Path
import openpyxl


def clean(value: object) -> str:
    return " ".join(str("" if value is None else value).replace("\xa0", " ").split())


def main(registry_path: str, workbook_path: str, output_path: str) -> None:
    registry = list(csv.DictReader(Path(registry_path).open(encoding="utf-8-sig"), delimiter="\t"))
    required = {"look_id", "excel_sheet", "excel_look_number"}
    if not registry or not required.issubset(registry[0]):
        raise ValueError("Registry must contain look_id, excel_sheet, and excel_look_number.")
    workbook = openpyxl.load_workbook(workbook_path, data_only=True, read_only=True)
    rows: list[dict[str, str]] = []
    seen: set[str] = set()
    for record in registry:
        look_id = clean(record["look_id"])
        sheet_name = clean(record["excel_sheet"])
        look_number = clean(record["excel_look_number"])
        if look_id in seen or not look_id.startswith("LOOK_") or sheet_name not in workbook.sheetnames:
            raise ValueError(f"Invalid registry mapping for {look_id}.")
        seen.add(look_id)
        sheet = workbook[sheet_name]
        start = None
        for row in range(1, sheet.max_row + 1):
            if clean(sheet.cell(row, 2).value) == look_number:
                start = row
                break
        if start is None:
            raise ValueError(f"{look_id}: look {look_number} not found in sheet {sheet_name}.")
        count = 0
        for row in range(start, sheet.max_row + 1):
            if row > start and clean(sheet.cell(row, 2).value):
                break
            fields = [clean(sheet.cell(row, column).value) for column in (3, 4, 5, 6)]
            if not any(fields):
                if count:
                    break
                continue
            if not all(fields):
                raise ValueError(f"{look_id}: incomplete product row {row}.")
            try:
                price = int(float(fields[2]))
            except ValueError as error:
                raise ValueError(f"{look_id}: invalid price at row {row}.") from error
            rows.append({"look_id": look_id, "type": fields[0], "brand": fields[1], "price": f"{price:,}".replace(",", " ") + " ₽", "article": fields[3]})
            count += 1
        if not count:
            raise ValueError(f"{look_id}: no product rows found.")
    output = Path(output_path)
    with output.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=["look_id", "type", "brand", "price", "article"], delimiter="\t")
        writer.writeheader()
        writer.writerows(rows)
    print(f"Wrote {len(rows)} product rows for {len(seen)} looks to {output}")


if __name__ == "__main__":
    raise SystemExit("BLOCKED: use build_verified_caption_data.py <caption-map.tsv> <source.xlsx> <caption-data.tsv> --provenance <caption-provenance.json>.")
